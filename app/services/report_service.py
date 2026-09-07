import io
from datetime import date, datetime
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy import desc

from app.models.attendance import AttendanceRecord
from app.models.employee import Employee

STATUS_LABELS = {
    "ON_TIME": "Đúng giờ",
    "LATE": "Đi muộn",
    "EARLY_LEAVE": "Về sớm",
    "LATE_AND_EARLY": "Muộn & Về sớm",
    "COMPLETED": "Hoàn thành",
    "INCOMPLETE": "Chưa check-out"
}

class ReportService:
    @staticmethod
    def generate_excel_report(
        db: Session,
        from_date: Optional[date] = None,
        to_date: Optional[date] = None,
        department: Optional[str] = None,
        employee_id: Optional[int] = None
    ) -> io.BytesIO:
        """
        Tạo báo cáo chấm công định dạng Excel chuyên nghiệp
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bảng Chấm Công"

        # Đặt font chung
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
        
        title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
        sub_font = Font(name="Calibri", size=10, italic=True)

        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        # Tiêu đề báo cáo
        ws.merge_cells("A1:K1")
        ws["A1"] = "BẢNG TỔNG HỢP CHẤM CÔNG NHÂN VIÊN (FACE ID)"
        ws["A1"].font = title_font
        ws["A1"].alignment = center_align

        ws.merge_cells("A2:K2")
        date_range_str = f"Từ ngày: {from_date.strftime('%d/%m/%Y') if from_date else 'Tất cả'} đến ngày: {to_date.strftime('%d/%m/%Y') if to_date else 'Hiện tại'}"
        ws["A2"] = f"{date_range_str} | Xuất lúc: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        ws["A2"].font = sub_font
        ws["A2"].alignment = center_align

        # Header Columns
        headers = [
            "STT", "Mã NV", "Họ và Tên", "Phòng Ban", "Ngày",
            "Giờ Vào (In)", "Giờ Ra (Out)", "Số Giờ Làm", "Đi Muộn (phút)", "Về Sớm (phút)", "Trạng Thái"
        ]

        row_idx = 4
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # Query dữ liệu
        query = db.query(AttendanceRecord).join(Employee)
        if from_date:
            query = query.filter(AttendanceRecord.record_date >= from_date)
        if to_date:
            query = query.filter(AttendanceRecord.record_date <= to_date)
        if department and department != "ALL":
            query = query.filter(Employee.department == department)
        if employee_id:
            query = query.filter(AttendanceRecord.employee_id == employee_id)

        records = query.order_by(desc(AttendanceRecord.record_date), Employee.emp_code).all()

        row_idx = 5
        for idx, rec in enumerate(records, 1):
            emp = rec.employee
            check_in = rec.check_in_time.strftime("%H:%M:%S") if rec.check_in_time else "--:--"
            check_out = rec.check_out_time.strftime("%H:%M:%S") if rec.check_out_time else "--:--"
            status_text = STATUS_LABELS.get(rec.status, rec.status)

            ws.cell(row=row_idx, column=1, value=idx).alignment = center_align
            ws.cell(row=row_idx, column=2, value=emp.emp_code).alignment = center_align
            ws.cell(row=row_idx, column=3, value=emp.full_name).alignment = left_align
            ws.cell(row=row_idx, column=4, value=emp.department).alignment = left_align
            ws.cell(row=row_idx, column=5, value=rec.record_date.strftime("%d/%m/%Y")).alignment = center_align
            ws.cell(row=row_idx, column=6, value=check_in).alignment = center_align
            ws.cell(row=row_idx, column=7, value=check_out).alignment = center_align
            ws.cell(row=row_idx, column=8, value=rec.work_hours).alignment = right_align
            ws.cell(row=row_idx, column=9, value=rec.late_minutes).alignment = right_align
            ws.cell(row=row_idx, column=10, value=rec.early_minutes).alignment = right_align
            
            status_cell = ws.cell(row=row_idx, column=11, value=status_text)
            status_cell.alignment = center_align

            # Màu sắc theo trạng thái
            if rec.status == "LATE" or rec.late_minutes > 0:
                status_cell.font = Font(name="Calibri", color="DC2626", bold=True) # Red
            elif rec.status in ["ON_TIME", "COMPLETED"]:
                status_cell.font = Font(name="Calibri", color="16A34A", bold=True) # Green
            elif rec.status == "INCOMPLETE":
                status_cell.font = Font(name="Calibri", color="CA8A04") # Yellow

            for c in range(1, 12):
                ws.cell(row=row_idx, column=c).border = thin_border

            row_idx += 1

        # Tự động căn chỉnh độ rộng cột
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row < 4:
                    continue
                val = str(cell.value or "")
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

report_service = ReportService()
